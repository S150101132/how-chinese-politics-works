"""Create an editable reading edition from prepare-book.mjs's checked tokens.

Usage: the bundled Python runtime, from the project root. Run the artifact
operation marker before the first creation, then render and inspect the output.
"""
import json
import re
from pathlib import Path
from html import unescape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

ROOT = Path.cwd()
BOOK = json.loads((ROOT / 'tmp/book/book.json').read_text())
OUT = ROOT / '交付'
OUT.mkdir(exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.78)
sec.left_margin = sec.right_margin = Inches(.8)
sec.header_distance = sec.footer_distance = Inches(.35)

def font_style(style, size, family='Noto Serif CJK SC'):
    style.font.name = family
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.append(fonts)
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + key), family)

font_style(doc.styles['Normal'], 12)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = Pt(20)
normal.space_after = Pt(6)
normal.widow_control = True
for name, size in [('Title', 26), ('Heading 1', 20), ('Heading 2', 15), ('Heading 3', 13)]:
    font_style(doc.styles[name], size, 'Noto Sans CJK SC')
    pf = doc.styles[name].paragraph_format
    pf.space_before = Pt(15); pf.space_after = Pt(9)
    pf.keep_with_next = True
    pf.line_spacing = Pt(size * 1.4)
font_style(doc.styles['Caption'], 10)
for name in ['List Bullet', 'List Number']:
    font_style(doc.styles[name], 12)
    doc.styles[name].paragraph_format.line_spacing = Pt(20)

# The bundled default template contains theme-colored paragraph rules.
# Remove them so the book's titles and prose use only typographic hierarchy.
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)

bookmarks = 0
current = None
ref_seen = set()

def bookmark(p, name):
    global bookmarks
    bookmarks += 1
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(bookmarks)); start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(bookmarks))
    p._p.append(start); p._p.append(end)

def hyperlink(p, text, url=None, anchor=None, superscript=False, bold=False, italic=False, small=False):
    h = OxmlElement('w:hyperlink')
    if url:
        h.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    else:
        h.set(qn('w:anchor'), anchor)
    r = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '31554B'); props.append(color)
    if superscript:
        vert = OxmlElement('w:vertAlign'); vert.set(qn('w:val'), 'superscript'); props.append(vert)
    if bold: props.append(OxmlElement('w:b'))
    if italic: props.append(OxmlElement('w:i'))
    if small or superscript:
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '20' if small else '17'); props.append(sz)
    r.append(props); t = OxmlElement('w:t'); t.text = text; r.append(t); h.append(r); p._p.append(h)

def plain(p, text, bold=False, italic=False, small=False):
    parts = re.split(r'(\[\^[^\]]+\])', unescape(text))
    for part in parts:
        if not part: continue
        m = re.fullmatch(r'\[\^([^\]]+)\]', part)
        if m:
            key = f"{current['id']}_note_{m[1]}"
            if key not in ref_seen:
                bookmark(p, key + '_ref'); ref_seen.add(key)
            hyperlink(p, f'[{m[1]}]', anchor=key, superscript=True)
        else:
            r = p.add_run(part); r.bold = bold; r.italic = italic
            if small: r.font.size = Pt(10)

def inlines(p, tokens, bold=False, italic=False, small=False):
    for t in tokens or []:
        kind = t['type']
        if kind == 'strong': inlines(p, t.get('tokens'), True, italic, small)
        elif kind == 'em': inlines(p, t.get('tokens'), bold, True, small)
        elif kind == 'link':
            hyperlink(p, unescape(t.get('text','')), url=t['href'], bold=bold, italic=italic, small=small)
        elif kind == 'br': p.add_run().add_break()
        elif kind == 'image': picture(t['href'], t.get('text',''))
        elif kind == 'html':
            if re.fullmatch(r'<br\s*/?>', t.get('text','')): p.add_run().add_break()
            else: plain(p, re.sub('<[^>]+>', '', t.get('text','')), bold, italic, small)
        elif 'tokens' in t: inlines(p, t['tokens'], bold, italic, small)
        else: plain(p, t.get('text', t.get('raw','')), bold, italic, small)

def picture(source, alt):
    pth = Path(source)
    if not pth.is_absolute(): pth = ROOT / pth
    if not pth.exists(): raise ValueError('Missing image '+str(pth))
    w, h = Image.open(pth).size
    width = min(6.7, 7.6 * w / h)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Exact body line height clips inline pictures in Writer and Word.
    # Auto line spacing expands this paragraph to the full picture height.
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(7)
    shape = p.add_run().add_picture(str(pth), width=Inches(width))
    shape._inline.docPr.set('descr', alt)

def table(t):
    rows = [t['header']] + t['rows']
    tab = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tab.autofit = False
    tab.alignment = 1
    cols = len(rows[0])
    widths = [6.7 / cols] * cols
    if cols == 2: widths = [2.05, 4.65]
    if cols == 3: widths = [1.4, 2.65, 2.65]
    for col, width in zip(tab.columns, widths): col.width = Inches(width)
    for i, row in enumerate(rows):
        trpr = tab.rows[i]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0: trpr.append(OxmlElement('w:tblHeader'))
        for j, data in enumerate(row):
            cell = tab.cell(i,j); cell.width = Inches(widths[j])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.line_spacing = Pt(15)
            inlines(p, data.get('tokens', []), bold=(i==0), small=True)
            if i == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:fill'),'EDF0EC'); cell._tc.get_or_add_tcPr().append(sh)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e = OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'C8CEC9'); borders.append(e)
    tab._tbl.tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def blocks(tokens):
    for t in tokens:
        kind = t['type']
        if kind in ['space', 'def']: continue
        if kind == 'heading':
            p = doc.add_paragraph(style='Heading '+str(min(t['depth'],3)))
            inlines(p,t.get('tokens'))
        elif kind in ['paragraph','text']:
            # A stand-alone image gets its own picture paragraph.
            if len(t.get('tokens',[])) == 1 and t['tokens'][0]['type'] == 'image':
                it=t['tokens'][0]; picture(it['href'],it.get('text',''))
            else:
                p=doc.add_paragraph()
                if re.fullmatch(r'\*\*[^*\n]{1,30}\*\*[:：]?',t.get('text','').strip()):
                    p.paragraph_format.keep_with_next=True
                inlines(p,t.get('tokens',[{'type':'text','text':t.get('text','')}]))
        elif kind == 'table': table(t)
        elif kind == 'figure': picture(t['path'],t['alt'])
        elif kind == 'code':
            if t.get('lang')=='mermaid': raise ValueError('Unrendered Mermaid')
            p=doc.add_paragraph(); p.paragraph_format.line_spacing=Pt(16)
            r=p.add_run(t['text']); r.font.name='Noto Sans CJK SC'; r.font.size=Pt(10.5)
        elif kind == 'blockquote': blocks(t['tokens'])
        elif kind == 'list':
            start=t.get('start',1) or 1
            for idx,item in enumerate(t['items']):
                parts=item.get('tokens',[])
                p=doc.add_paragraph(style='List Number' if t['ordered'] else 'List Bullet')
                # Word list numbers are editable; each new list restarts through explicit text.
                if t['ordered']:
                    p.style='Normal'; p.paragraph_format.left_indent=Inches(.25)
                    p.paragraph_format.first_line_indent=Inches(-.25); p.add_run(str(start+idx)+'. ')
                if parts and parts[0]['type'] in ['text','paragraph']:
                    inlines(p,parts[0].get('tokens',[{'type':'text','text':parts[0].get('text','')}]))
                    blocks(parts[1:])
                else: inlines(p,[{'type':'text','text':item.get('text','')}])
        elif kind == 'hr': doc.add_paragraph()
        else: raise ValueError('Unknown block type '+kind)

header=sec.header.paragraphs[0]
header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=header.add_run('中国政治是怎么运转的'); r.font.size=Pt(9)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
for text in ['PAGE','NUMPAGES']:
    if text=='NUMPAGES': footer.add_run(' / ')
    f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),text); footer._p.append(f)

p=doc.add_paragraph('中国政治是怎么运转的',style='Title')
p.paragraph_format.space_before=Pt(100)
doc.add_paragraph('从制度、人物与现实事件读懂当代中国')
doc.add_paragraph('第一版 1.0')
doc.add_paragraph('现实资料截止 2026年9月29日')
doc.add_paragraph('本书通过真实人物和事件解释机构、程序与制度变化，帮助普通读者读懂新闻中的动作，并分清事实、解释和未知。')
doc.add_paragraph('注释保留在各章末尾；正文中的编号可跳转到依据，注释中的“返回”可跳回首次引用。')
doc.add_page_break()
p=doc.add_paragraph('目录',style='Heading 1'); bookmark(p,'book_contents')
for c in BOOK:
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3)
    hyperlink(p,c['title'],anchor=c['id'])

for c in BOOK:
    current=c
    p=doc.add_paragraph(c['title'],style='Heading 1')
    p.paragraph_format.page_break_before=True
    bookmark(p,c['id'])
    blocks(c['tokens'])
    if c['notes']:
        doc.add_paragraph('注释与资料',style='Heading 2')
        for k, ts in c['notesTokens'].items():
            p=doc.add_paragraph(); p.paragraph_format.line_spacing=Pt(15)
            p.paragraph_format.space_after=Pt(7)
            key=f"{c['id']}_note_{k}"
            bookmark(p,key); plain(p,f'[{k}] ',small=True); inlines(p,ts,small=True)
            p.add_run(' '); hyperlink(p,'返回',anchor=key+'_ref',small=True)

doc.core_properties.title='中国政治是怎么运转的'
doc.core_properties.subject='从制度、人物与现实事件读懂当代中国'
doc.core_properties.author=''
doc.core_properties.last_modified_by=''
doc.core_properties.language='zh-CN'
doc.core_properties.comments='现实资料截止2026年9月29日；制作方式见附录F。'
output=OUT/'中国政治是怎么运转的.docx'
doc.save(output)
print(json.dumps({'output':str(output),'sections':len(BOOK),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'images':len(doc.inline_shapes)},ensure_ascii=False))
